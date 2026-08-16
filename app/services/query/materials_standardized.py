# -*- coding: utf-8 -*-
"""已规整物资台账只读筛选（optv3/material-standization-filtering.md）。

数据来自 fact_inventory（+ fact_asset 工器具）LEFT JOIN dim_material。
物资编码只用正式 material_code，不用 material_id 兜底。
筛选参数一律占位符绑定，排序字段走白名单。
"""
from __future__ import annotations

import json
from datetime import datetime
from io import BytesIO
from typing import Any
from urllib.parse import quote

from fastapi import HTTPException
from fastapi.responses import Response

from app import config
from app.repositories import biz_conn

STANDARD_CATEGORIES = (
    "维护材料",
    "低值易耗品",
    "备品备件",
    "公用工器具",
    "个人工器具",
)

SORT_WHITELIST = {
    "material_code": "material_code",
    "material_name": "material_name",
    "stock_qty": "stock_qty",
}

EXPORT_COLUMNS = [
    ("material_code", "物资编码"),
    ("material_name", "物资名称"),
    ("category", "物资种类"),
    ("location", "存放区域"),
    ("spec", "规格型号"),
    ("unit", "单位"),
    ("stock_qty", "库存数量"),
    ("status", "状态"),
]

_CATS_SQL = ", ".join("'" + c.replace("'", "''") + "'" for c in STANDARD_CATEGORIES)

# 物资种类优先取标准枚举（source_sheet / category），否则回退业务类别。
# 物资编码：仅正式编码；与内部 ID 相同或空则视为未维护。
LEDGER_SQL = f"""
SELECT
  CASE
    WHEN m.material_code IS NOT NULL
         AND trim(CAST(m.material_code AS VARCHAR)) <> ''
         AND trim(CAST(m.material_code AS VARCHAR)) <> trim(CAST(i.material_id AS VARCHAR))
    THEN trim(CAST(m.material_code AS VARCHAR))
    ELSE NULL
  END AS material_code,
  COALESCE(m.material_name, '') AS material_name,
  CASE
    WHEN i.source_sheet IN ({_CATS_SQL}) THEN i.source_sheet
    WHEN COALESCE(i.category, m.category) IN ({_CATS_SQL}) THEN COALESCE(i.category, m.category)
    ELSE COALESCE(
      NULLIF(trim(CAST(i.category AS VARCHAR)), ''),
      NULLIF(trim(CAST(i.source_sheet AS VARCHAR)), ''),
      '未分类'
    )
  END AS category,
  NULLIF(trim(CAST(i.location AS VARCHAR)), '') AS location,
  m.spec AS spec,
  COALESCE(NULLIF(trim(CAST(m.unit AS VARCHAR)), ''), i.unit) AS unit,
  i.stock_qty AS stock_qty,
  CASE
    WHEN i.stock_qty IS NULL THEN '未维护'
    WHEN i.stock_qty > 0 THEN '在库'
    ELSE '无库存'
  END AS status,
  i.source_release_id AS source_release_id,
  i.row_key AS row_key,
  i.source_file AS source_file,
  i.material_id AS material_id
FROM fact_inventory i
LEFT JOIN dim_material m ON i.material_id = m.material_id
UNION ALL
SELECT
  CASE
    WHEN a.material_code IS NOT NULL
         AND trim(CAST(a.material_code AS VARCHAR)) <> ''
    THEN trim(CAST(a.material_code AS VARCHAR))
    ELSE NULL
  END AS material_code,
  COALESCE(a.asset_name, '') AS material_name,
  CASE
    WHEN a.source_sheet IN ({_CATS_SQL}) THEN a.source_sheet
    WHEN a.source_sheet IS NOT NULL AND a.source_sheet LIKE '%个人%' THEN '个人工器具'
    ELSE '公用工器具'
  END AS category,
  NULLIF(trim(CAST(a.location AS VARCHAR)), '') AS location,
  CAST(NULL AS VARCHAR) AS spec,
  a.unit AS unit,
  a.asset_qty AS stock_qty,
  COALESCE(NULLIF(trim(CAST(a.status AS VARCHAR)), ''), '—') AS status,
  a.source_release_id AS source_release_id,
  a.row_key AS row_key,
  a.source_file AS source_file,
  a.asset_code AS material_id
FROM fact_asset a
"""


def _split_csv(raw: str | None) -> list[str]:
    if not raw:
        return []
    seen: list[str] = []
    for part in str(raw).split(","):
        v = part.strip()
        if v and v not in seen:
            seen.append(v)
    return seen


def parse_categories(raw: str | None) -> list[str]:
    return _split_csv(raw)


def parse_locations(raw: str | None) -> list[str]:
    return _split_csv(raw)


def parse_keyword(raw: str | None) -> str:
    q = (raw or "").strip()
    if len(q) > 100:
        q = q[:100]
    return q


def _where(categories: list[str], locations: list[str], q: str) -> tuple[str, list[Any]]:
    clauses: list[str] = []
    params: list[Any] = []
    if categories:
        placeholders = ", ".join("?" * len(categories))
        clauses.append(f"category IN ({placeholders})")
        params.extend(categories)
    if locations:
        placeholders = ", ".join("?" * len(locations))
        clauses.append(f"location IN ({placeholders})")
        params.extend(locations)
    if q:
        # contains() 按字面匹配，避免 % / _ 被当成 SQL LIKE 通配符扫出全表。
        # 关键字只查业务字段：物资编码、物资名称。不按内部 material_id 检索。
        clauses.append(
            "("
            " (material_code IS NOT NULL AND contains(lower(CAST(material_code AS VARCHAR)), lower(?)))"
            " OR (material_name IS NOT NULL AND contains(lower(CAST(material_name AS VARCHAR)), lower(?)))"
            ")"
        )
        params.extend([q, q])
    if not clauses:
        return "", params
    return " WHERE " + " AND ".join(clauses), params


def _sort_sql(sort_by: str | None, sort_order: str | None) -> str:
    col = SORT_WHITELIST.get((sort_by or "").strip(), "material_name")
    order = "DESC" if str(sort_order or "").strip().lower() == "desc" else "ASC"
    return f" ORDER BY {col} {order} NULLS LAST, material_name ASC, location ASC"


def _null_if_nan(v: Any) -> Any:
    if v is None:
        return None
    try:
        if v != v:  # NaN
            return None
    except Exception:
        pass
    return v


def _item_from_row(row: dict[str, Any]) -> dict[str, Any]:
    code = _null_if_nan(row.get("material_code"))
    if code is not None:
        code = str(code).strip() or None
    return {
        "material_code": code,
        "material_id": _null_if_nan(row.get("material_id")),
        "material_name": _null_if_nan(row.get("material_name")) or "",
        "category": _null_if_nan(row.get("category")) or "",
        "location": _null_if_nan(row.get("location")),
        "spec": _null_if_nan(row.get("spec")),
        "unit": _null_if_nan(row.get("unit")),
        "stock_qty": _null_if_nan(row.get("stock_qty")),
        "status": _null_if_nan(row.get("status")) or "—",
        "source_file": _null_if_nan(row.get("source_file")),
        "source_release_id": _null_if_nan(row.get("source_release_id")),
        "row_key": _null_if_nan(row.get("row_key")),
    }


def _fetch_df(con, sql: str, params: list[Any]):
    return con.execute(sql, params).fetchdf() if params else con.execute(sql).fetchdf()


def _fetch_one(con, sql: str, params: list[Any]):
    cur = con.execute(sql, params) if params else con.execute(sql)
    return cur.fetchone()


def list_filters() -> dict[str, list[str]]:
    con = biz_conn()
    try:
        loc_df = con.execute(
            f"SELECT DISTINCT location FROM ({LEDGER_SQL}) t "
            "WHERE location IS NOT NULL AND trim(CAST(location AS VARCHAR)) <> '' "
            "ORDER BY location"
        ).fetchdf()
        cat_df = con.execute(
            f"SELECT DISTINCT category FROM ({LEDGER_SQL}) t "
            "WHERE category IS NOT NULL AND trim(CAST(category AS VARCHAR)) <> '' "
            "ORDER BY category"
        ).fetchdf()
    finally:
        con.close()
    locations = [
        str(v)
        for v in (loc_df["location"].tolist() if len(loc_df) else [])
        if v is not None and str(v).strip() not in ("", "nan", "None", "<NA>")
    ]
    extras = [
        str(v)
        for v in (cat_df["category"].tolist() if len(cat_df) else [])
        if v is not None
        and str(v).strip() not in ("", "nan", "None", "<NA>")
        and str(v) not in STANDARD_CATEGORIES
    ]
    return {"categories": list(STANDARD_CATEGORIES) + extras, "locations": locations}


def list_standardized(
    *,
    categories: str | None = None,
    locations: str | None = None,
    q: str | None = None,
    limit: int = 20,
    offset: int = 0,
    sort_by: str | None = None,
    sort_order: str | None = None,
) -> dict[str, Any]:
    cats = parse_categories(categories)
    locs = parse_locations(locations)
    keyword = parse_keyword(q)
    limit = max(1, min(int(limit or 20), 500))
    offset = max(0, int(offset or 0))
    where_sql, params = _where(cats, locs, keyword)
    order_sql = _sort_sql(sort_by, sort_order)
    base = f"SELECT * FROM ({LEDGER_SQL}) ledger{where_sql}"
    con = biz_conn()
    try:
        total_row = _fetch_one(con, f"SELECT COUNT(*) FROM ({base}) x", params)
        total = int(total_row[0]) if total_row else 0
        page_sql = f"{base}{order_sql} LIMIT ? OFFSET ?"
        df = _fetch_df(con, page_sql, params + [limit, offset])
        sum_sql = (
            f"SELECT category, COUNT(*) AS n FROM ({base}) x GROUP BY category ORDER BY category"
        )
        sum_df = _fetch_df(con, sum_sql, params)
    finally:
        con.close()
    raw_items = json.loads(df.to_json(orient="records")) if len(df) else []
    items = [_item_from_row(r) for r in raw_items]
    counts = {
        str(r["category"]): int(r["n"])
        for r in (json.loads(sum_df.to_json(orient="records")) if len(sum_df) else [])
    }
    by_category = [{"category": c, "count": counts.get(c, 0)} for c in STANDARD_CATEGORIES]
    extra = [
        {"category": k, "count": v}
        for k, v in sorted(counts.items())
        if k not in set(STANDARD_CATEGORIES)
    ]
    return {
        "items": items,
        "total": total,
        "summary": {"by_category": by_category + extra},
        "filters": {"categories": cats, "locations": locs, "q": keyword},
        "limit": limit,
        "offset": offset,
    }


def export_standardized(
    *,
    categories: str | None = None,
    locations: str | None = None,
    q: str | None = None,
    sort_by: str | None = None,
    sort_order: str | None = None,
) -> Response:
    cats = parse_categories(categories)
    locs = parse_locations(locations)
    keyword = parse_keyword(q)
    where_sql, params = _where(cats, locs, keyword)
    order_sql = _sort_sql(sort_by, sort_order)
    cap = max(1, int(config.EXPORT_ROW_LIMIT))
    sql = f"SELECT * FROM ({LEDGER_SQL}) ledger{where_sql}{order_sql} LIMIT ?"
    con = biz_conn()
    try:
        total_row = _fetch_one(
            con, f"SELECT COUNT(*) FROM (SELECT * FROM ({LEDGER_SQL}) ledger{where_sql}) x", params
        )
        total = int(total_row[0]) if total_row else 0
        if total <= 0:
            raise HTTPException(
                status_code=400,
                detail={
                    "code": "EMPTY_EXPORT",
                    "message": "当前筛选条件查询结果为空，不支持导出，请重新设置筛选条件",
                },
            )
        df = _fetch_df(con, sql, params + [cap])
    finally:
        con.close()
    out = df[[c[0] for c in EXPORT_COLUMNS]].copy()

    def _export_text(v: Any, empty: str = "") -> str:
        if v is None:
            return empty
        try:
            import pandas as pd

            if pd.isna(v):
                return empty
        except Exception:
            pass
        try:
            if v != v:
                return empty
        except Exception:
            pass
        s = str(v).strip()
        if s in ("", "None", "nan", "<NA>"):
            return empty
        return s

    out["material_code"] = out["material_code"].map(lambda v: _export_text(v, "未维护"))
    for col in ("material_name", "category", "location", "spec", "unit", "status"):
        out[col] = out[col].map(lambda v: _export_text(v, ""))
    out.columns = [c[1] for c in EXPORT_COLUMNS]

    def _xlsx_cell(v: Any) -> Any:
        if v is None:
            return None
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            try:
                if v != v:
                    return None
            except Exception:
                pass
            return v
        s = str(v)
        if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
            return "'" + v
        if s[:1] in ("=", "+", "-", "@"):
            return "'" + s
        return v

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "筛选结果"
    headers = [c[1] for c in EXPORT_COLUMNS]
    ws.append(headers)
    for row in out.itertuples(index=False, name=None):
        ws.append([_xlsx_cell(v) for v in row])
    buf = BytesIO()
    wb.save(buf)
    filename = f"物资台账_筛选结果_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                'attachment; filename="material_filter.xlsx"; '
                f"filename*=UTF-8''{quote(filename)}"
            )
        },
    )
